"""Build app/templates/jssd_minutes.docx — the minutes as a FILL-IN Word template (docxtpl), like the
service-letter template the user gave as the reference (2026-09-20).

The layout lives in the Word file; the service only supplies values (app/minutes_template.py). Every value
is a named slot — {{ venue }}, {%p for a in attendees %} — and a slot with no value prints "xxx...xxx".
Open the result in Word and change anything: wording, fonts, spacing, fixed lines. Slots are listed in
templates/README.md.

Built from docx_export's own building blocks, so page, margins, tabs, spacing and keep-together settings
are the ones tools/check_jssd_layout.py already holds to the manual. Run from the repo root:

    docker run --rm -e PYTHONPATH=/app -w /app -v "$PWD/app:/app" -v "$PWD/tools:/tools:ro" \\
        --entrypoint python mom-prompt-service:local /tools/make_minutes_template.py /app/templates/jssd_minutes.docx
"""
import sys

import docx
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Emu, Inches, Pt

from docx_export import (CENTRE, COL, LEFT, SIGNATURE_FEEDS, TAB, TEXT_W, _blank, _centre, _field, _fmt, _keep,
                         _page_setup, _row, _run, _table)


def tag(container, text):
    """A paragraph holding only a {%p … %} tag: docxtpl removes it when it renders."""
    container.add_paragraph().add_run(text)


def tag_row(t, widths, text):
    """A table row holding only a {%tr … %} tag in its first cell: removed when rendered."""
    _row(t, widths).cells[0].paragraphs[0].add_run(text)


def spacer(t, widths):
    """The blank row above each entry — two-line spacing between entries (Part 1 para 7) — kept with the
    rows below it for the first entry of a list, so a heading never ends a page alone."""
    tag_row(t, widths, "{%tr if loop.first %}")
    _keep(_row(t, widths))
    tag_row(t, widths, "{%tr else %}")
    _row(t, widths)
    tag_row(t, widths, "{%tr endif %}")


def numbered(p, number_slot, text_runs, keep=False, right=None):
    """"N.<tab>text" — number at the margin, text after a 0.5 in tab, later lines back at the margin."""
    _fmt(p, tabs=(TAB,), keep=keep)
    if right is not None:
        p.paragraph_format.right_indent = right
    p.add_run(f"{{{{ {number_slot} }}}}.\t")
    for text, bold in text_runs:
        _run(p, text, bold=bold)
    return p


def sub(p, number_slot, text_slot, indent=TAB, right=None):
    """"N.n.<tab>text" — under the paragraph text, later lines under the number."""
    _fmt(p, left=indent, tabs=(Emu(indent + TAB),))
    if right is not None:
        p.paragraph_format.right_indent = right
    p.add_run(f"{{{{ {number_slot} }}}}\t{{{{ {text_slot} }}}}")
    return p


def build(path):
    doc = docx.Document()
    _page_setup(doc, draft=False)

    # ── superscription (AD; Part 1 paras 19-20) ────────────────────────────────────────────────────
    tag(doc, "{%p if draft %}")
    _centre(doc.add_paragraph(), "DRAFT")
    _blank(doc)
    tag(doc, "{%p endif %}")
    p = _fmt(doc.add_paragraph(), align=LEFT)
    p.add_run("Tele: {{ telephone }}")
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    _run(p, "{{ precedence }}", bold=True)
    tag(doc, "{%p if copy_no %}")
    _blank(doc)
    p = _fmt(doc.add_paragraph(), align=LEFT)
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t{{ copy_no }}")
    tag(doc, "{%p endif %}")
    _blank(doc)
    tag(doc, "{%p for line in address %}")
    _fmt(doc.add_paragraph(), align=LEFT).add_run("{{ line }}")
    tag(doc, "{%p endfor %}")
    _blank(doc)
    _fmt(doc.add_paragraph(), align=LEFT).add_run("{{ file_reference }}\tdt {{ date_of_issue }}")
    _blank(doc)

    # ── title: place, time, date and purpose (AD note 2) ───────────────────────────────────────────
    _run(_fmt(doc.add_paragraph(), align=CENTRE, keep=True), "{{ title }}", bold=True)
    _blank(doc, keep=True)

    # ── 1. The following were present:- (AD note 3) ────────────────────────────────────────────────
    numbered(doc.add_paragraph(), "present_number", [("The following were present:-", False)], keep=True)
    widths = (TEXT_W - TAB - Inches(2.9), Inches(1.9), COL)
    t = _table(doc, widths, indent=TAB)
    tag_row(t, widths, "{%tr for a in attendees %}")
    spacer(t, widths)                               # one blank line above each name, kept with the first
    cells = _row(t, widths).cells
    sub(cells[0].paragraphs[0], "a.number", "a.name", indent=0)
    _fmt(cells[1].paragraphs[0], align=LEFT).add_run("{{ a.appointment }}")
    _fmt(cells[2].paragraphs[0], align=LEFT).add_run("{{ a.label }}")
    tag_row(t, widths, "{%tr endfor %}")

    # ── INTRODUCTION (Ch 6 para 16.12) ─────────────────────────────────────────────────────────────
    tag(doc, "{%p if introduction %}")
    _blank(doc, keep=True)
    _centre(doc.add_paragraph(), "INTRODUCTION", keep=True)
    tag(doc, "{%p for b in introduction %}")
    _blank(doc, keep=True)
    tag(doc, "{%p if b.subs %}")
    numbered(doc.add_paragraph(), "b.number", [("{{ b.text }}", False)], keep=True)
    tag(doc, "{%p else %}")
    numbered(doc.add_paragraph(), "b.number", [("{{ b.text }}", False)])
    tag(doc, "{%p endif %}")
    tag(doc, "{%p for s in b.subs %}")
    tag(doc, "{%p if loop.first %}")
    _blank(doc, keep=True)                          # a page never ends on a lead-in line (Part 1 para 36)
    tag(doc, "{%p else %}")
    _blank(doc)
    tag(doc, "{%p endif %}")
    sub(doc.add_paragraph(), "s.number", "s.text")
    tag(doc, "{%p endfor %}")
    tag(doc, "{%p endfor %}")
    tag(doc, "{%p endif %}")

    # ── ITEM I, II, III… with Action and Info columns (AD notes 4-6; Ch 6 paras 16.7-16.16) ─────────
    tag(doc, "{%p if items %}")
    _blank(doc)
    widths = (TEXT_W - 2 * COL, COL, COL)
    t = _table(doc, widths)
    head = _keep(_row(t, widths, header=True)).cells
    for c, label in zip(head[1:], ("Action", "Info")):
        _centre(c.paragraphs[0], label)
    tag_row(t, widths, "{%tr for item in items %}")
    _keep(_row(t, widths))
    _centre(_keep(_row(t, widths)).cells[0].paragraphs[0], "ITEM {{ item.roman }} – {{ item.title }}",
            underline=True, keep=True)
    _centre(_keep(_row(t, widths)).cells[0].paragraphs[0], "({{ item.classification }})", bold=False, keep=True)
    tag_row(t, widths, "{%tr for e in item.entries %}")
    spacer(t, widths)                               # the first line of an item stays with its heading
    # A figures lead-in paragraph stays with its first figure.
    tag_row(t, widths, "{%tr if e.figures %}")
    for keep in (True, False):
        cells = _row(t, widths).cells
        numbered(cells[0].paragraphs[0], "e.number",
                 [("{% if e.decision %}Decision.{% endif %}", True), ("{% if e.decision %} {% endif %}{{ e.text }}", False)],
                 keep=keep, right=Inches(0.1))
        _centre(cells[1].paragraphs[0], "{{ e.action }}", bold=False)
        _centre(cells[2].paragraphs[0], "{{ e.info }}", bold=False)
        tag_row(t, widths, "{%tr else %}" if keep else "{%tr endif %}")
    tag_row(t, widths, "{%tr for f in e.figures %}")
    spacer(t, widths)
    sub(_row(t, widths).cells[0].paragraphs[0], "f.number", "f.text", right=Inches(0.1))
    tag_row(t, widths, "{%tr endfor %}")
    tag_row(t, widths, "{%tr endfor %}")
    tag_row(t, widths, "{%tr endfor %}")
    tag(doc, "{%p endif %}")

    # ── closing and signature block, nine line feeds down (AD; Ch 6 paras 15, 16.17) ────────────────
    _blank(doc)
    numbered(doc.add_paragraph(), "closing_number",
             [("Agreement with the minutes will be assumed unless amendments are received by {{ amendments_by }}.", False)],
             keep=True)
    for _ in range(SIGNATURE_FEEDS - 1):
        _blank(doc, keep=True)
    for i, line in enumerate(("({{ secretary_name }})", "{{ secretary_rank }}", "Secretary")):
        _fmt(doc.add_paragraph(), align=LEFT, keep=i < 2).add_run(line)

    # ── distribution (AD; Part 1 paras 67-69) ───────────────────────────────────────────────────────
    _blank(doc, keep=True)
    widths = (Inches(2.4), Inches(1.4), Inches(1.1), TEXT_W - Inches(4.9))
    t = _table(doc, widths)
    head = _keep(_row(t, widths)).cells
    for c, label in zip(head, ("Distribution", "No of Copies", "Copy No", "Remarks")):
        _run(_fmt(c.paragraphs[0], align=LEFT), label, bold=True)
    tag_row(t, widths, "{%tr for d in distribution %}")
    _row(t, widths)                                 # one blank line above each row
    for c, slot in zip(_row(t, widths).cells, ("addressee", "copies", "copy_no", "remarks")):
        _fmt(c.paragraphs[0], align=LEFT).add_run(f"{{{{ d.{slot} }}}}")
    tag_row(t, widths, "{%tr endfor %}")

    # ── classification head and foot of every page; page count on page 1; page numbers after ────────
    # (Part 1 paras 12.1, 16, 18; App B notes 27-28). Unclassified: no marking at all (para 11.5).
    normal = doc.styles["Normal"]
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True

    def para(part, first=False):
        p = part.paragraphs[0] if first else part.add_paragraph()
        p.style = normal
        return p

    first = sec.first_page_header
    para(first, True).add_run("{%p if classification %}")
    _centre(para(first), "{{ classification }}")
    para(first).add_run("{%p if pages_shown %}")
    p = _fmt(para(first), align=CENTRE)
    n = ([" NUMPAGES "], "{{ page_count }}")
    words = ([" NUMPAGES \\* CardText \\* FirstCap "], "{{ page_count_word }}")
    inner = ([" IF ", n, ' < 10 "(', words, ' pages)" "(', n, ' pages)" '], "{{ page_count_text }}")
    _field(p, [" IF ", n, ' = 1 "(Only page)" "', inner, '" '], "{{ page_count_text }}")
    para(first).add_run("{%p endif %}")
    para(first)                                                  # the text starts two lines below
    para(first).add_run("{%p else %}")
    para(first).paragraph_format.line_spacing = Pt(1)            # nothing above the first line of text
    para(first).add_run("{%p endif %}")

    rest = sec.header
    para(rest, True).add_run("{%p if classification %}")
    _centre(para(rest), "{{ classification }}")
    para(rest)
    para(rest).add_run("{%p endif %}")
    _field(_fmt(para(rest), align=CENTRE), [" PAGE "], "2")
    para(rest)

    for footer in (sec.first_page_footer, sec.footer):
        para(footer, True).add_run("{%p if classification %}")
        para(footer)
        para(footer).add_run("{%p endif %}")
        _centre(para(footer), "{{ classification }}")

    doc.save(path)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "jssd_minutes.docx")
    print("written", sys.argv[1] if len(sys.argv) > 1 else "jssd_minutes.docx")
