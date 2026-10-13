"""One dictionary of header values: the screen's fields and buttons, the user's prompt, earlier messages and the
HQ's template — the screen wins — and "....xxxx....xxxx...." wherever nobody gave a value (2026-10-13)."""
import io

import docx
import header_values as hv
from docx_export import MISSING, build_mom_docx
from kafka_contract import parse_job

ok = fail = 0


def check(name, cond, extra=""):
    global ok, fail
    if cond:
        ok += 1
        print("  ok  ", name)
    else:
        fail += 1
        print("  FAIL", name, extra)


MOM = {"title": "Weekly maintenance review", "summary": "The meeting covered the water supply.",
       "attendees": [{"name": "Maj Rohit Negi", "role": "Station Commander"}], "agenda": ["Water supply"],
       "key_points": ["42 families in Block C have had no water since 1900 hrs yesterday."],
       "decisions": ["The line will be repaired by 24 September."], "action_items": [], "key_figures": []}
PROMPT = ("Prepare the minutes of the weekly maintenance review.\n"
          "Venue: Station HQ Conference Room\n"
          "Tele: 0194-2450101\n"
          "File ref: 1234/56/Trg, dt: 12 Oct 26\n"
          "Secretary: Maj Rohit Negi\n"
          "Address: HQ 7 Inf Bde; C/O 56 APO; PIN 900111\n"
          "Distribution: HQ 7 Inf Bde | 2 | for action; File\n"
          "Classification: restricted")


def job(prompt="", meta=None, **echo):
    return parse_job({"tenant_id": "t", "conversationId": "c1", "prompt": prompt, "mom_meta": meta or {}, **echo})


print("1. the prompt only")
meta, where, rest = hv.collect(job(PROMPT), {})
check("venue", meta.get("venue") == "Station HQ Conference Room", meta.get("venue"))
check("telephone", meta.get("telephone") == "0194-2450101", meta.get("telephone"))
check("file reference and its date, written on one line", meta.get("file_ref") == "1234/56/Trg"
      and meta.get("issue_date") == "12 Oct 26", (meta.get("file_ref"), meta.get("issue_date")))
check("secretary", (meta.get("secretary") or {}).get("name") == "Maj Rohit Negi", meta.get("secretary"))
check("address, line by line", meta.get("address") == ["HQ 7 Inf Bde", "C/O 56 APO", "PIN 900111"], meta.get("address"))
check("distribution rows", meta.get("distribution") == [{"addressee": "HQ 7 Inf Bde", "copies": "2", "remarks": "for action"},
                                                        {"addressee": "File", "copies": "", "remarks": ""}], meta.get("distribution"))
check("classification as a grade word", meta.get("classification") == "RESTRICTED", meta.get("classification"))
check("every one of them came from the prompt", set(where.values()) == {"the prompt"}, where)
check("the header lines leave the text the minutes are written from",
      rest == "Prepare the minutes of the weekly maintenance review.", rest)
meta2, _, _ = hv.collect(job("Tele: 9654396200, Venue: Room B"), {})
check("several on one line", meta2.get("telephone") == "9654396200" and meta2.get("venue") == "Room B", meta2)
meta3, _, rest3 = hv.collect(job("We met in Conference Room B and the telephone rang twice."), {})
check("prose with no labels gives nothing", meta3 == {} and "Conference Room B" in rest3, meta3)
meta4, _, _ = hv.collect(job("Classification: as applicable"), {})
check("a classification that names no grade is not set", "classification" not in meta4, meta4)
meta5, _, _ = hv.collect(job("Venue: " + "x" * 400), {})
check("a value too long to be a detail is left alone", "venue" not in meta5, meta5)

print("2. the screen only")
screen = {"venue": "IAF Head Quarters", "meeting_date": "22 Sep 26", "meeting_time": "1030",
          "secretary": {"name": "Sqn Ldr A Mehta", "rank": "Sqn Ldr"}}
meta, where, rest = hv.collect(job("Prepare the minutes of the weekly review.", screen), {})
check("the screen's fields are taken", meta.get("venue") == "IAF Head Quarters" and meta.get("meeting_time") == "1030", meta)
check("marked as the screen's", where.get("venue") == "the screen", where)
check("the prompt is untouched", rest == "Prepare the minutes of the weekly review.", rest)
meta, where, _ = hv.collect(job("Prepare the minutes.", {}, metaData='{"file_ref": "SPEC/9901/Trg", "venue": "Room 4"}',
                                grading="SECRET"), {})
check("header values inside metaData count as the screen's",
      meta.get("file_ref") == "SPEC/9901/Trg" and where.get("file_ref") == "the screen", (meta, where))
check("the grading button sets the classification", meta.get("classification") == "SECRET", meta)
meta, _, _ = hv.collect(job("Prepare the minutes.", {}, grading="normal"), {})
check("a grading that names no grade sets nothing", "classification" not in meta, meta)

print("3. both — the screen wins")
meta, where, rest = hv.collect(job(PROMPT, {"venue": "IAF Head Quarters", "classification": "CONFIDENTIAL"}), {})
check("venue: the screen's, not the prompt's", meta.get("venue") == "IAF Head Quarters"
      and where.get("venue") == "the screen", (meta.get("venue"), where.get("venue")))
check("classification: the screen's", meta.get("classification") == "CONFIDENTIAL", meta.get("classification"))
check("what only the prompt gave is still taken", meta.get("telephone") == "0194-2450101"
      and where.get("telephone") == "the prompt", where)
check("the prompt's header lines still leave the source text", "Tele:" not in rest and "Venue:" not in rest, rest)
meta, where, _ = hv.collect(job("Secretary: Maj Rohit Negi", {"secretary": {"rank": "Maj"}}), {})
check("name from the prompt, rank from the screen", meta["secretary"] == {"rank": "Maj", "name": "Maj Rohit Negi"}
      and where["secretary_rank"] == "the screen" and where["secretary_name"] == "the prompt", (meta, where))

print("4. earlier messages, then the template, then nothing")
earlier = {"amendments_by": "30 Sep 26", "venue": "Old Room"}
template = {"telephone": "0194-9999999", "address": ["HQ 15 Corps"], "file_ref": "TPL/1"}
meta, where, _ = hv.collect(job(PROMPT), template, earlier)
check("an earlier message fills what nobody else gave", meta.get("amendments_by") == "30 Sep 26"
      and where.get("amendments_by") == "an earlier message", where)
check("the prompt still beats an earlier message", meta.get("venue") == "Station HQ Conference Room", meta.get("venue"))
check("the template fills only what is left", where.get("telephone") == "the prompt" and "address" in meta
      and where.get("address") == "the prompt", where)
meta, where, _ = hv.collect(job("Prepare the minutes."), template)
check("with no screen and no prompt, the template is used", meta.get("telephone") == "0194-9999999"
      and where.get("telephone") == "the template", where)
check("precedence, copy number and the rest stay unset", "precedence" not in meta and "meeting_date" not in meta, meta)

print("5. in the Word file")


def text_of(meta):
    d = docx.Document(io.BytesIO(build_mom_docx(MOM, meta)))
    return "\n".join([p.text for p in d.paragraphs]
                     + [c.text for t in d.tables for row in t.rows for c in row.cells])


full, _, _ = hv.collect(job(PROMPT, {"meeting_date": "22 Sep 26", "meeting_time": "1030"}), {})
body = text_of(full)
check('the placeholder is "....xxxx....xxxx...."', MISSING == "....xxxx....xxxx....", MISSING)
check("values given are printed", "0194-2450101" in body and "STATION HQ CONFERENCE ROOM" in body
      and "1234/56/Trg" in body and "HQ 7 Inf Bde" in body, body[:400])
check("the classification is printed", "RESTRICTED" in body, body[:200])
check("what nobody gave shows the placeholder", MISSING in body, body[:400])
empty = text_of({})
for label, piece in (("telephone", "Tele: " + MISSING), ("file reference and date", MISSING + "\tdt " + MISSING),
                     ("the title's venue", "HELD AT " + MISSING), ("secretary", "(" + MISSING + ")")):
    check(f"with nothing given, {label} shows the placeholder", piece in empty, empty[:500])
check("the placeholder's dots are not trimmed in the title", MISSING + "\n" in empty + "\n" or MISSING in empty)

print(f"\n{ok} passed, {fail} failed")
