"""`documents` + `questionnaire` in the job, and the header by priority: the form, then the prompt, then the documents
(2026-10-14). Parsing, question → field, the checks on the model's answers, the merge, the template slots, a whole
job and its follow-ups. The model, MinIO and Elasticsearch are stand-ins."""
import copy
import io
import json
import sys
import types

sys.modules.setdefault("pypdfium2", types.ModuleType("pypdfium2"))     # only outside the image: not used here

import agenda_items
import header_fields as H
import minutes
import minutes_edit as me
import minutes_template
from kafka_contract import build_ack, parse_job

ok = fail = 0


def check(name, cond, got=None):
    global ok, fail
    if cond:
        ok += 1
    else:
        fail += 1
        print(f"FAIL {name}: {got!r}")


# ── the input, as the curl sends it ─────────────────────────────────────────────────────────────────────────
body = {
    "conversationId": "uyhbfyghjn", "tenant_id": "rfved", "prompt": "Prepare the minutes of the meeting.",
    "path": "AsItIs", "metaData": "{}", "user": False,
    "documents": [{"fid": "ikjhbvfghj", "miniourl": "corr-files-upload/livetest-ask-feb0b19043.docx",
                   "fileName": "ikjhvdfghnmws"}],
    "questionnaire": [{"question": "file reference number", "answer": "AMCID12323"}],
}
job = parse_job(body)
check("documents → file_urls", job.file_urls == ["corr-files-upload/livetest-ask-feb0b19043.docx"], job.file_urls)
check("documents → names", job.document_names == ["ikjhvdfghnmws"], job.document_names)
check("questionnaire parsed", job.questionnaire == {"file reference number": "AMCID12323"}, job.questionnaire)
check("path AsItIs not a file", "AsItIs" not in job.file_urls)
ack = build_ack(job, success=True, description="x", bucket="b", object_key="k")
check("inputs not echoed", "documents" not in ack and "questionnaire" not in ack, sorted(ack))
check("other fields echoed", ack.get("path") == "AsItIs" and ack.get("metaData") == "{}" and ack.get("user") is False)

# dict form, typo in question, fid-only document, file_urls kept first
j2 = parse_job({"conversationId": "c", "file_urls": ["mom/a.pdf"], "document_names": ["a.pdf"],
                "documents": [{"fid": "6aad24087fb955220d1ad0cf", "fileName": "t1.docx"},
                              {"fid": "x", "miniourl": "mom/b.docx", "fileName": "b.docx"},
                              {"fid": "y", "miniourl": "mom/a.pdf", "fileName": "dup"}],
                "questionnaire": {"File Refernce Number": "A/1", "Telephone number": 1334675, "Blank": ""}})
check("file_urls then documents", j2.file_urls == ["mom/a.pdf", "mom/b.docx"], j2.file_urls)
check("names aligned", j2.document_names == ["a.pdf", "b.docx"], j2.document_names)
check("fid-only → file_fids", j2.file_fids == ["6aad24087fb955220d1ad0cf"], j2.file_fids)
check("dict questionnaire, blank skipped, number as text",
      j2.questionnaire == {"File Refernce Number": "A/1", "Telephone number": "1334675"}, j2.questionnaire)
j3 = parse_job({"conversationId": "c", "questionnaire": json.dumps([{"question": "venue", "answer": "Room B"}])})
check("questionnaire sent as JSON text", j3.questionnaire == {"venue": "Room B"}, j3.questionnaire)
j4 = parse_job({"conversationId": "c", "path": "mom/x.pdf"})
check("no documents: path still a file", j4.file_urls == ["mom/x.pdf"], j4.file_urls)

# ── question → field ────────────────────────────────────────────────────────────────────────────────────────
for q, f in [("file reference number", "file_ref"), ("File Refernce Number", "file_ref"), ("File No.", "file_ref"),
             ("Telephone number", "telephone"), ("Tele", "telephone"), ("Venue", "venue"), ("Date", "meeting_date"),
             ("Subject", "title"), ("Secretary", "secretary_name"), ("Rank of the Secretary", "secretary_rank"),
             ("Distribution list", "distribution"), ("Date of issue", "issue_date"),
             ("Name of the officer", ""), ("Remarks", ""), ("Budget head", "")]:
    check(f"field_of({q!r})", H.field_of(q) == f, H.field_of(q))
check("slot name", H.slot_name("File Reference No.") == "file_reference_no", H.slot_name("File Reference No."))

# ── 1. the job's own values: questionnaire over mom_meta ────────────────────────────────────────────────────
j5 = parse_job({"conversationId": "c", "prompt": "p",
                "mom_meta": {"file_ref": "META/1", "venue": "Hall A", "classification": "RESTRICTED"},
                "questionnaire": [{"question": "file reference number", "answer": "AMCID12323"},
                                  {"question": "Telephone number", "answer": "1334675"},
                                  {"question": "Address", "answer": "HQ 7 Inf Bde, C/O 56 APO, PIN 900111"},
                                  {"question": "Secretary", "answer": "Rohit Negi"},
                                  {"question": "Subject", "answer": "Annual training review"},
                                  {"question": "Officer in charge", "answer": "Col A K Singh"}]})
given, title = H.inputs(j5)
check("questionnaire beats mom_meta", given["file_ref"] == "AMCID12323", given.get("file_ref"))
check("mom_meta kept where the form is silent", given["venue"] == "Hall A" and given["classification"] == "RESTRICTED")
check("telephone", given["telephone"] == "1334675")
check("address split", given["address"] == ["HQ 7 Inf Bde", "C/O 56 APO", "PIN 900111"], given.get("address"))
check("secretary", given["secretary"] == {"name": "Rohit Negi"}, given.get("secretary"))
check("title", title == "Annual training review", title)
check("every answer kept by its question", given["fields"]["Officer in charge"] == "Col A K Singh"
      and given["fields"]["file reference number"] == "AMCID12323", given.get("fields"))
check("job's mom_meta untouched", j5.mom_meta == {"file_ref": "META/1", "venue": "Hall A", "classification": "RESTRICTED"})

# ── 2./3. prompt and documents, model answers checked ───────────────────────────────────────────────────────
prompt = ("Minutes of the meeting held at Conference Room B on 12th September 2026 at 1030 hrs. File No A/123/MoM. "
          "Tele: 2450101. Mark it top secret.")
reply = {"venue": "Conference Room B", "meeting_date": "12 September 2026", "meeting_time": "1030 hrs",
         "file_ref": "A/123/MoM", "telephone": "Tele: 2450101", "classification": "SECRET",
         "secretary_name": "Maj Verma", "address": "Telephone number here"}
wanted = list(H.FIELDS)
got = H._checked("prompt", reply, prompt, wanted, "[T]")
check("prompt: venue kept", got.get("venue") == "Conference Room B", got)
check("prompt: date kept (ordinal, month written out)", got.get("meeting_date") == "12 September 2026", got)
check("prompt: file ref kept", got.get("file_ref") == "A/123/MoM", got)
check("prompt: Tele prefix dropped", got.get("telephone") == "2450101", got)
check("prompt: 'top secret' never read as SECRET", "classification" not in got, got)
check("prompt: invented secretary refused", "secretary" not in got, got)
check("prompt: placeholder refused", "address" not in got, got)
got2 = H._checked("prompt", {"classification": "restricted"}, "please mark it restricted", ["classification"], "[T]")
check("prompt: the one grade named is taken", got2.get("classification") == "RESTRICTED", got2)
check("number inside a bigger number is not 'written'", not H.written_in("123", "ref 41234 x"))

# the model call itself is stubbed: read() asks only for fields not already given, never raises
asked = {}
def fake_ask(kind, text, wanted):
    asked[kind] = list(wanted)
    return {"telephone": "999", "file_ref": "DOC/9"} if kind == "documents" else {}
real_ask = H._ask_model
H._ask_model = fake_ask
docs_text = "Office of the HQ. File No DOC/9. Tele 999."
found = H.read("documents", docs_text, given, "[T]")
check("documents: only their fields, only the missing ones",
      asked["documents"] == ["secretary_rank", "distribution"], asked.get("documents"))
check("documents: nothing asked for is not taken", found == {}, found)
H._ask_model = lambda *a: (_ for _ in ()).throw(RuntimeError("model down"))
check("a failed call costs nothing", H.read("prompt", "some prompt", {}, "[T]") == {})
H._ask_model = real_ask
check("no text → no call", H.read("prompt", "  ", {}, "[T]") == {})

# ── the merge: form > prompt > earlier messages > documents > template ──────────────────────────────────────
top = {"file_ref": "FORM/1", "secretary": {"rank": "Maj"}}
meta, used = H.merge(top, [
    ("prompt", {"file_ref": "PROMPT/1", "telephone": "111", "secretary": {"name": "Rohit Negi"}}),
    ("earlier messages", {"telephone": "222", "venue": "Room C", "fields": {"a": "old", "b": "kept"}}),
    ("documents", {"telephone": "333", "address": ["Doc line"], "secretary": {"name": "X", "rank": "Col"}}),
    ("template", {"address": ["Tpl line"], "distribution": [{"addressee": "Comdt"}], "telephone": "444"}),
])
check("form wins", meta["file_ref"] == "FORM/1", meta)
check("prompt beats earlier messages and documents", meta["telephone"] == "111", meta)
check("earlier messages fill what prompt left", meta["venue"] == "Room C", meta)
check("documents beat template", meta["address"] == ["Doc line"], meta)
check("template last", meta["distribution"] == [{"addressee": "Comdt"}], meta)
check("secretary: rank from form, name from prompt, nothing from the document",
      meta["secretary"] == {"rank": "Maj", "name": "Rohit Negi"}, meta.get("secretary"))
check("which layer gave what", used == {"prompt": ["telephone", "secretary"], "earlier messages": ["venue"],
                                        "documents": ["address"], "template": ["distribution"]}, used)
check("merge leaves its input alone", top == {"file_ref": "FORM/1", "secretary": {"rank": "Maj"}})

# ── the template: built-in slots, form names, custom questions ──────────────────────────────────────────────
mom = {"title": "Annual training review", "summary": "The meeting reviewed training.", "agenda": ["Training"],
       "key_points": ["Training was reviewed."], "decisions": [], "action_items": [], "attendees": []}
ctx = minutes_template.context(mom, given)
check("file_reference slot", ctx["file_reference"] == "AMCID12323", ctx["file_reference"])
check("form-named slot", ctx["file_reference_number"] == "AMCID12323" and ctx["telephone_number"] == "1334675")
check("subject slot", ctx["subject"] == "Annual training review", ctx.get("subject"))
check("custom question slot", ctx["officer_in_charge"] == "Col A K Singh", ctx.get("officer_in_charge"))
check("built-in slot not replaced", not isinstance(ctx["title"], str))
check("unknown slot → marker in an empty header", minutes_template.context(mom, {})["file_reference_number"] == "xxx...xxx")

from docx import Document
d = Document()
d.add_paragraph("Ref {{ file_reference_number }} Tele {{ telephone }} Subj {{ subject }} OIC {{ officer_in_charge }} "
                "Other {{ never_asked }}")
buf = io.BytesIO(); d.save(buf)
out = minutes_template._fill(buf.getvalue(), ctx, "", "", False)
text = " ".join(p.text for p in out.paragraphs)
check("client template filled", text == "Ref AMCID12323 Tele 1334675 Subj Annual training review OIC Col A K Singh "
      "Other xxx...xxx", text)

doc = minutes_template._fill(open(minutes_template.DEFAULT_TEMPLATE, "rb").read(), ctx, "", "", False)
allt = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
check("default template: file ref", "AMCID12323" in allt)
check("default template: telephone", "1334675" in allt)
check("default template: address", "C/O 56 APO" in allt)
check("default template: title", "ANNUAL TRAINING REVIEW" in allt)

# ── an edit carrying the questionnaire ──────────────────────────────────────────────────────────────────────
ej = parse_job({"conversationId": "c", "mode": "edit",
                "questionnaire": [{"question": "file reference number", "answer": "NEW/2"},
                                  {"question": "Subject", "answer": "Annual training review"},
                                  {"question": "Officer in charge", "answer": "Col B"}]})
m2, meta2 = copy.deepcopy(mom), copy.deepcopy(given)
done = me._from_form(ej, m2, meta2)
check("edit: changed values reported", any("file reference is now NEW/2 (was AMCID12323)" in d for d in done)
      and any("Officer in charge is now Col B (was Col A K Singh)" in d for d in done), done)
check("edit: unchanged title not reported", not any("title" in d for d in done), done)
check("edit: applied", meta2["file_ref"] == "NEW/2" and meta2["fields"]["Officer in charge"] == "Col B")
check("edit: same form again changes nothing", me._from_form(ej, m2, meta2) == [])


# ── a whole job, then follow-ups ───────────────────────────────────────────────────────────────────────────
class Store:
    def __init__(self):
        self.files, self.puts = {}, []

    def download(self, url):
        for key in (url, url.split("/", 1)[-1]):
            if key in self.files:
                return self.files[key]
        raise RuntimeError("NoSuchKey")

    def upload(self, key, data, mime):
        self.files[key] = data
        self.puts.append(key)
        return "mom", key


class Index:
    def index_mom(self, *a, **k):
        pass


DOC = ("MINUTES. File No DOC/77/MoM. Tele 0194-111. HQ 9 Sector, C/O 99 APO. Col Rao opened the meeting. "
       "Water supply in Block C was discussed; the line will be repaired by 24 September. "
       "Convoy: all 12 vehicles to be checked by 1600 hrs.")
WRITER_META = []


class LLM:
    def generate(self, source, meta):
        WRITER_META.append(meta)
        return {"content": {
            "header": {"title": "Weekly review", "venue": "Hall from the document"},
            "summary": "The meeting addressed the water line repair in Block C.",
            "attendees": [{"name": "Col Rao", "role": "Chairman"}],
            "agenda": ["Water supply", "Convoy"],
            "key_points": ["The water line in Block C was discussed."],
            "decisions": ["The line will be repaired by 24 September."],
            "action_items": [], "key_figures": []}}


class C:
    def __init__(self):
        self.store, self.index, self.chunks, self.llm = Store(), Index(), Index(), LLM()


READ = []
minutes._read_documents = lambda job, c: (READ.append(list(job.file_urls)) or [("notes.docx", DOC)])
RENDERED = []
minutes.build_mom_docx = lambda mom, meta, template=None: (RENDERED.append((mom, meta)) or
                                                           json.dumps([mom, meta], sort_keys=True).encode())
me.build_mom_docx = minutes.build_mom_docx
agenda_items.group = lambda *a, **k: None
minutes.agenda_items.group = agenda_items.group


def fake_model(kind, text, wanted):
    if kind == "prompt":     # the prompt states a file reference, a telephone and a venue
        return {"file_ref": "PROMPT/5", "telephone": "2450101", "venue": "Conference Room B", "meeting_time": ""}
    return {"file_ref": "DOC/77/MoM", "telephone": "0194-111", "address": "HQ 9 Sector ; C/O 99 APO",
            "secretary_name": "Maj Ghost"}            # Maj Ghost is not in the document: refused


H._ask_model = fake_model
c = C()
body = {"conversationId": "q-001", "tenant_id": "t", "path": "AsItIs", "queryId": "q1",
        "prompt": "Minutes of the meeting held at Conference Room B. File No PROMPT/5, Tele 2450101.",
        "documents": [{"fid": "ikjhbvfghj", "miniourl": "corr-files-upload/livetest-ask-feb0b19043.docx",
                       "fileName": "notes.docx"}],
        "questionnaire": [{"question": "file reference number", "answer": "AMCID12323"},
                          {"question": "Officer in charge", "answer": "Col A K Singh"}]}
ack = minutes.process(parse_job(body), c)
check("job succeeds", ack["message"] == "SUCCESS", ack)
check("document read from its miniourl", READ[-1] == ["corr-files-upload/livetest-ask-feb0b19043.docx"], READ)
mom, meta = RENDERED[-1]
check("file ref: form beats prompt and document", meta.get("file_ref") == "AMCID12323", meta)
check("telephone: prompt beats document", meta.get("telephone") == "2450101", meta)
check("venue: prompt beats the writer's", meta.get("venue") == "Conference Room B", meta)
check("address: from the document", meta.get("address") == ["HQ 9 Sector", "C/O 99 APO"], meta)
check("invented secretary refused", "secretary" not in meta, meta)
check("custom answer kept for the template", meta.get("fields", {}).get("Officer in charge") == "Col A K Singh", meta)
check("ack echoes, no inputs", ack.get("queryId") == "q1" and "questionnaire" not in ack and "documents" not in ack)
check("file named after the document", "/MoM-notes-" in ack["summaryObjectKey"], ack["summaryObjectKey"])

# a follow-up with only a new questionnaire answer, same document: an edit, no rewrite
writes = len(WRITER_META)
ack2 = minutes.process(parse_job({**body, "prompt": "", "questionnaire": [
    {"question": "file reference number", "answer": "NEW/9"}]}), c)
check("form-only follow-up succeeds", ack2["message"] == "SUCCESS", ack2)
check("form-only follow-up is an edit, not a rewrite", len(WRITER_META) == writes, len(WRITER_META))
check("form-only follow-up says what changed", "file reference is now NEW/9 (was AMCID12323)" in ack2["description"],
      ack2["description"])
_, meta2 = RENDERED[-1]
check("edit keeps the rest of the header", meta2.get("telephone") == "2450101" and meta2.get("file_ref") == "NEW/9", meta2)

# the form beats the chat in the same message
me._ask = lambda mom, meta, instr, earlier="", refused="": {"changes": [
    {"op": "set_meta", "list": "", "index": -1, "quote": "", "field": "file_ref", "value": "CHAT/1", "role": "",
     "owner": "", "due": ""}]}
ack3 = minutes.process(parse_job({**body, "prompt": "change the file reference to CHAT/1",
                                  "questionnaire": [{"question": "file reference number", "answer": "FORM/3"}]}), c)
_, meta3 = RENDERED[-1]
check("form beats chat in an edit", meta3.get("file_ref") == "FORM/3", (meta3.get("file_ref"), ack3["description"]))

# a question with a form change: answered AND changed, not written again
me._ask = lambda mom, meta, instr, earlier="", refused="": {"changes": [
    {"op": "answer", "list": "", "index": -1, "quote": "", "field": "", "value": "The agenda was water and convoy.",
     "role": "", "owner": "", "due": ""}]}
writes = len(WRITER_META)
ack4 = minutes.process(parse_job({**body, "prompt": "what was the agenda?",
                                  "questionnaire": [{"question": "Telephone number", "answer": "777"}]}), c)
check("question + form: not rewritten", len(WRITER_META) == writes)
check("question + form: changed and answered", "telephone is now 777" in ack4["description"]
      and "water and convoy" in ack4["description"], ack4["description"])

# a new job with no questionnaire: prompt, then documents, as before
c2 = C()
ack5 = minutes.process(parse_job({**body, "conversationId": "q-002", "questionnaire": []}), c2)
_, meta5 = RENDERED[-1]
check("no form: prompt's file ref", meta5.get("file_ref") == "PROMPT/5", meta5)
check("writer told the form's values only", WRITER_META[-1] == {}, WRITER_META[-1])

print(f"{ok} passed, {fail} failed")
